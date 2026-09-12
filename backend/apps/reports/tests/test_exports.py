"""
Phase 7: downloads.

The browser never receives a whole report as JSON — it asks for a file, and the file is written
row by row straight from the database cursor.
"""

import csv
import io

import pytest
from openpyxl import load_workbook

from .conftest import REPORTS

pytestmark = pytest.mark.django_db

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def download(client, key: str, export_format: str = "csv", **params):
    query = "&".join(
        f"{name}={value}" for name, value in {"format": export_format, **params}.items()
    )
    # A download link clicked in the browser sends this Accept header, not application/json.
    return client.get(f"{REPORTS}{key}/export/?{query}", HTTP_ACCEPT="text/html")


def content(response) -> bytes:
    if response.streaming:
        return b"".join(response.streaming_content)
    return response.content


def csv_rows(response) -> list[list[str]]:
    text = content(response).decode("utf-8-sig")
    return [row for row in csv.reader(io.StringIO(text)) if row]


class TestCSV:
    def test_it_downloads_as_a_named_file(self, admin_api, ledger):
        response = download(admin_api, "bookings")

        assert response.status_code == 200
        assert response["Content-Type"].startswith("text/csv")
        assert "attachment; filename=" in response["Content-Disposition"]
        assert "yathra-bookings" in response["Content-Disposition"]

    def test_it_carries_the_header_and_every_row(self, admin_api, ledger):
        rows = csv_rows(download(admin_api, "bookings", range="today"))

        assert rows[0][:3] == ["Reference", "Booked on", "Customer"]
        assert len(rows) == 1 + 3
        assert {row[2] for row in rows[1:]} == {
            "Alice Perera",
            "Bob Silva",
            "Carol Jayasuriya",
        }

    def test_money_is_written_plainly_for_a_spreadsheet(self, admin_api, ledger):
        rows = csv_rows(download(admin_api, "revenue", range="today"))
        header, day = rows[0], rows[1]

        assert header == ["Date", "Bookings", "Payments", "Gross", "Refunds", "Net"]
        assert day[3:] == ["8700.00", "2500.00", "6200.00"]

    def test_it_streams_rather_than_building_the_whole_body(self, admin_api, ledger):
        assert download(admin_api, "passengers", range="all").streaming is True


class TestExcel:
    def test_it_is_a_real_workbook(self, admin_api, ledger):
        response = download(admin_api, "payments", "xlsx", range="today")

        assert response.status_code == 200
        assert response["Content-Type"] == XLSX_TYPE
        assert ".xlsx" in response["Content-Disposition"]

        sheet = load_workbook(io.BytesIO(content(response)), read_only=True).active
        rows = list(sheet.values)
        assert rows[0][0] == "Transaction"
        assert len(rows) == 1 + 3

    def test_the_header_row_is_bold(self, admin_api, ledger):
        response = download(admin_api, "routes", "xlsx", range="all")
        sheet = load_workbook(io.BytesIO(content(response))).active

        assert sheet["A1"].font.bold is True


class TestPDF:
    def test_it_is_a_pdf(self, admin_api, ledger):
        response = download(admin_api, "occupancy", "pdf", range="all")

        assert response.status_code == 200
        assert response["Content-Type"] == "application/pdf"
        assert content(response).startswith(b"%PDF-")

    def test_every_report_can_be_printed(self, admin_api, ledger):
        for key in ("bookings", "passengers", "revenue", "routes", "cancellations", "payments"):
            response = download(admin_api, key, "pdf", range="all")
            assert response.status_code == 200, key
            assert content(response).startswith(b"%PDF-"), key


class TestBadRequests:
    def test_an_unknown_format_is_refused(self, admin_api, ledger):
        response = download(admin_api, "bookings", "docx")

        assert response.status_code == 400
        assert response["Content-Type"].startswith("application/json")
        assert "format" in response.json()["error"]["details"]

    def test_an_unknown_report_is_a_404(self, admin_api, ledger):
        assert download(admin_api, "profits").status_code == 404

    def test_a_customer_cannot_download_the_takings(self, api_client, customer, ledger):
        api_client.force_authenticate(customer)

        assert api_client.get(f"{REPORTS}revenue/export/").status_code == 403
